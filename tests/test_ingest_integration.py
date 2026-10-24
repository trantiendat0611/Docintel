"""Test chạy Docling thật trên ba loại file: PDF số, PDF scan, DOCX.

Đánh dấu ``slow`` vì phải tải model và OCR trên CPU mất vài phút.
Chạy bằng: uv run pytest -m slow
"""

from io import BytesIO

import pytest

from app.db.models import Block, DocStatus, Document, IngestProgress, PageStage
from app.ingestion.parse import parse_document
from app.ingestion.service import ingest, register_upload
from scripts.make_scans import pdf_to_scan

pytestmark = pytest.mark.slow


@pytest.fixture(scope="session")
def scan_pdf_bytes(sample_pdf_bytes) -> bytes:
    return pdf_to_scan(sample_pdf_bytes, dpi=200, seed=1)


@pytest.fixture(scope="session")
def docx_bytes() -> bytes:
    from docx import Document as DocxDocument

    document = DocxDocument()
    document.add_heading("Quy chế đào tạo", level=1)
    document.add_paragraph("Điều 1. Sinh viên được đăng ký tối đa 24 tín chỉ mỗi học kỳ.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Học kỳ"
    table.cell(0, 1).text = "Số tín chỉ"
    table.cell(1, 0).text = "Học kỳ 1"
    table.cell(1, 1).text = "24"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_parses_digital_pdf_with_table_and_provenance(sample_pdf_bytes):
    result = parse_document(sample_pdf_bytes, "hoa-don-001.pdf")

    assert result.ocr_pages == set()  # PDF số thì không chạy OCR
    assert result.page_count == 1
    assert [b.position for b in result.blocks] == list(range(len(result.blocks)))

    tables = [b for b in result.blocks if b.block_type == "table"]
    assert len(tables) == 1
    assert "Đơn giá" in tables[0].content and "|" in tables[0].content

    for block in result.blocks:
        assert block.page == 1
        assert block.bbox is not None
        assert block.bbox["page_width"] > 0
        assert not block.from_ocr


def test_ocr_runs_only_for_scanned_pages(scan_pdf_bytes):
    result = parse_document(scan_pdf_bytes, "hoa-don-001-scan.pdf")

    assert result.ocr_pages == {1}
    assert all(block.from_ocr for block in result.blocks)
    text = " ".join(block.content for block in result.blocks)
    # Không khẳng định OCR đúng hết dấu: đó là việc của phép đo CER ở thí nghiệm riêng.
    assert "HO" in text.upper()
    assert any(char.isdigit() for char in text)


def test_parses_docx(docx_bytes):
    result = parse_document(docx_bytes, "quy-che.docx")

    contents = " ".join(block.content for block in result.blocks)
    assert "tín chỉ" in contents
    assert any(block.block_type == "table" for block in result.blocks)


def test_ingest_writes_blocks_and_progress(session, storage, sample_pdf_bytes):
    document, created = register_upload(
        session, storage, filename="hoa-don-001.pdf", data=sample_pdf_bytes
    )
    assert created

    ingest(session, storage, document.id)
    session.refresh(document)

    assert document.status == DocStatus.READY.value
    assert document.page_count == 1
    assert document.ocr_page_count == 0
    assert storage.exists(document.parsed_key)

    blocks = session.query(Block).filter_by(document_id=document.id).all()
    assert len(blocks) >= 3

    progress = session.query(IngestProgress).filter_by(document_id=document.id).all()
    assert [p.stage for p in progress] == [PageStage.PARSED.value]


def test_ingest_is_rerunnable_without_duplicating_blocks(session, storage, sample_pdf_bytes):
    document, _ = register_upload(
        session, storage, filename="hoa-don-001.pdf", data=sample_pdf_bytes
    )
    ingest(session, storage, document.id)
    first = session.query(Block).filter_by(document_id=document.id).count()

    ingest(session, storage, document.id)
    assert session.query(Block).filter_by(document_id=document.id).count() == first


def test_ingest_marks_failure_for_corrupt_file(session, storage):
    document, _ = register_upload(
        session, storage, filename="hong.pdf", data=b"%PDF-1.7\nnoi dung hong"
    )
    with pytest.raises(Exception):  # noqa: B017 - Docling ném nhiều loại lỗi khác nhau
        ingest(session, storage, document.id)

    from app.ingestion.service import mark_failed

    mark_failed(session, document.id, "ConversionError: file hỏng")
    session.refresh(document)
    assert session.get(Document, document.id).status == DocStatus.FAILED.value
